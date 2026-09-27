"""Setul de referință etichetat manual: eșantionul de rânduri CNAM și fișierul Excel de etichetare.

Folosire (o singură dată, NU face parte din run_pipeline.py, pentru că etichetele sunt muncă manuală):
    python src/labeling.py
Creează data/labels/etichetare_cnam.xlsx (de completat manual) și data/labels/esantion_info.csv
(straturile și ponderile, pentru calculul metricilor). Nu suprascrie un fișier existent.

Eșantion stratificat: 100 de rânduri din cele potrivite exact după normalizare și 100 din cele
nepotrivite, cu seed fix. Ponderea fiecărui rând = mărimea stratului / mărimea eșantionului din strat,
astfel încât metricile calculate pe eșantion să poată fi raportate la întreaga listă CNAM.
Setul este folosit DOAR la evaluarea finală, nu la antrenarea sau ajustarea modelelor.
"""
from difflib import SequenceMatcher
from pathlib import Path

import numpy as np
import pandas as pd

from clean import key_text
from normalize import strip_accents

ROOT = Path(__file__).resolve().parents[1]
INTERIM = ROOT / "data" / "interim"
LABELS = ROOT / "data" / "labels"
SEED = 42
N_PER_STRATUM = 100
N_CANDIDATES = 10


def dci_key(s: pd.Series) -> pd.Series:
    """DCI comparabilă între surse: fără diacritice, fără spații în jurul lui „+”."""
    return strip_accents(key_text(s)).str.replace(r"\s*\+\s*", "+", regex=True)


def draw_sample(ro: pd.DataFrame, nom: pd.DataFrame) -> pd.DataFrame:
    """100 de rânduri potrivite exact + 100 nepotrivite, amestecate (ca stratul să nu fie vizibil)."""
    full = set(zip(nom.denumire_norm, nom.doza_norm, nom.forma_norm))
    ro = ro.assign(id_cnam=ro.index,
                   strat=["potrivit exact" if k in full else "nepotrivit"
                          for k in zip(ro.denumire_norm, ro.doza_norm, ro.forma_norm)])
    rng = np.random.default_rng(SEED)
    parts = []
    for strat, g in ro.groupby("strat"):
        idx = rng.choice(g.index, size=min(N_PER_STRATUM, len(g)), replace=False)
        parts.append(g.loc[idx].assign(pondere=len(g) / min(N_PER_STRATUM, len(g))))
    sample = pd.concat(parts)
    return sample.iloc[rng.permutation(len(sample))].reset_index(drop=True)


def candidates(row: pd.Series, nom: pd.DataFrame) -> pd.DataFrame:
    """Produsele din Nomenclator care ar putea fi perechea corectă: aceeași DCI sau denumire asemănătoare.
    Ordonate după asemănare (denumire, apoi doză, formă și divizare egale)."""
    a = row.denumire_norm if isinstance(row.denumire_norm, str) else ""
    name_sim = nom.denumire_norm.fillna("").map(lambda x: SequenceMatcher(None, a, x).ratio())
    pool = nom[(nom.dci_k == row.dci_k) | name_sim.isin(name_sim.nlargest(20))].copy()
    pool["scor"] = (0.6 * name_sim[pool.index] + 0.15 * (pool.doza_norm == row.doza_norm)
                    + 0.15 * (pool.forma_norm == row.forma_norm)
                    + 0.10 * (key_text(pool.divizare) == key_text(pd.Series([row.divizare])).iloc[0]))
    return pool.sort_values("scor", ascending=False).head(N_CANDIDATES)


INSTRUCTIUNI = [
    "Cum se completează fișierul",
    "",
    "Fiecare bloc începe cu un rând albastru: medicamentul din lista CNAM.",
    "Sub el sunt până la 10 produse din Nomenclator care ar putea fi același medicament (candidați).",
    "",
    "Întrebarea pentru fiecare bloc: care dintre candidați este EXACT același produs ca rândul CNAM?",
    "Același produs = aceeași denumire comercială, aceeași doză, aceeași formă farmaceutică și aceeași divizare (ambalaj).",
    "Diferențele de scriere nu contează (ex. „comp. film.” = „comprimate filmate”, „875 mg + 125 mg” = „875 mg/125 mg”).",
    "",
    "Ce completați (o singură variantă pentru fiecare bloc):",
    "A) Dacă găsiți produsul corect printre candidați: puneți „x” în coloana „alege (x)” pe rândul acelui candidat.",
    "B) Dacă produsul corect nu e printre candidați, dar îl găsiți pe nomenclator.amdm.gov.md: scrieți codul lui în coloana „cod manual” de pe rândul albastru.",
    "C) Dacă medicamentul nu există în Nomenclator: alegeți „NICIUNUL” în coloana „status” de pe rândul albastru.",
    "D) Dacă nu sunteți sigură: alegeți „INCERT” în coloana „status” și scrieți motivul în „comentariu”.",
    "",
    "Dacă mai mulți candidați par corecți (ex. aceeași denumire, doză, formă și divizare, dar deținători diferiți),",
    "alegeți-l pe cel al cărui deținător/producător corespunde cu producătorul din CNAM și notați situația în „comentariu”.",
    "",
    "Ordinea candidaților este dată de un calcul simplu de asemănare și poate greși; verificați fiecare candidat, nu doar primul.",
    "Nu modificați celelalte coloane și nu ștergeți rânduri.",
]


def write_excel(sample: pd.DataFrame, nom: pd.DataFrame, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "etichetare"
    head = ["id", "tip", "alege (x)", "status", "cod manual", "comentariu", "cod", "denumire", "DCI",
            "doză", "formă", "divizare", "volum", "deținător / producător", "țara"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    blue, yellow = PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="FFF2CC")
    top = Border(top=Side(style="medium"))
    dv_x = DataValidation(type="list", formula1='"x"', allow_blank=True)
    dv_status = DataValidation(type="list", formula1='"NICIUNUL,INCERT"', allow_blank=True)
    ws.add_data_validation(dv_x)
    ws.add_data_validation(dv_status)

    for _, r in sample.iterrows():
        ws.append([int(r.id_cnam), "CNAM", None, None, None, None, None, r.denumire, r.dci, r.doza,
                   r.forma, r.divizare, r.volum, r.producator, r.tara])
        i = ws.max_row
        for cell in ws[i]:
            cell.fill, cell.border = blue, top
            cell.font = Font(bold=True)
        for col in "DEF":
            ws[f"{col}{i}"].fill = yellow   # celulele de completat pe rândul CNAM
        dv_status.add(f"D{i}")
        for _, cnd in candidates(r, nom).iterrows():
            ws.append([int(r.id_cnam), "candidat", None, None, None, None, cnd.cod, cnd.denumire, cnd.dci,
                       cnd.doza, cnd.forma, cnd.divizare, cnd.volum, cnd.detinator, cnd.tara])
            ws[f"C{ws.max_row}"].fill = yellow
            dv_x.add(f"C{ws.max_row}")

    widths = [7, 9, 9, 11, 13, 25, 12, 34, 28, 18, 34, 12, 10, 28, 14]
    for col, w in zip("ABCDEFGHIJKLMNO", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column_letter in "FHIK")

    ins = wb.create_sheet("Instrucțiuni", 0)
    for line in INSTRUCTIUNI:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=14)
    ins.column_dimensions["A"].width = 140
    wb.active = 0
    wb.save(path)


def read_labels(path: Path = LABELS / "etichetare_cnam.xlsx") -> pd.DataFrame:
    """Citește etichetele completate: un rând per id_cnam, cu cod_corect (sau NaN dacă NICIUNUL) și status."""
    df = pd.read_excel(path, sheet_name="etichetare", dtype=str)
    cnam = df[df.tip == "CNAM"].set_index("id")
    chosen = df[(df.tip == "candidat") & (df["alege (x)"].str.strip().str.lower() == "x")].groupby("id").cod.agg(list)
    out = pd.DataFrame({"status": cnam["status"], "cod_manual": cnam["cod manual"], "ales": chosen.reindex(cnam.index)})
    out["cod_corect"] = out.cod_manual.where(out.cod_manual.notna(),
                                            out.ales.map(lambda x: x[0] if isinstance(x, list) and len(x) == 1 else np.nan))
    out["probleme"] = np.where(out.ales.map(lambda x: isinstance(x, list) and len(x) > 1), "mai multe x", "")
    out.loc[out.status.isna() & out.cod_corect.isna() & (out.probleme == ""), "probleme"] = "necompletat"
    out.index = out.index.astype(int)
    return out.drop(columns=["ales"])


JOURNAL = LABELS / "jurnal_etichetare.csv"
PREFILL_NOTE = "x pre-completat automat (candidat identic), de verificat"


def log_changes(rows: pd.DataFrame) -> None:
    """Adaugă în jurnal orice modificare a fișierului de etichetare făcută altfel decât manual
    (pre-completări automate, corecturi), ca setul de referință să rămână documentat."""
    cols = ["id_cnam", "tip", "inainte", "dupa", "motiv"]
    old = pd.read_csv(JOURNAL, dtype=str) if JOURNAL.exists() else pd.DataFrame(columns=cols)
    pd.concat([old, rows[cols].astype(str)]).to_csv(JOURNAL, index=False)


def prefill(path: Path = LABELS / "etichetare_cnam.xlsx") -> pd.DataFrame:
    """Pune „x” la candidatul identic cu rândul CNAM (denumire, doză, formă normalizate și divizare egale),
    doar dacă există EXACT unul și blocul nu a fost deja completat manual. Nu șterge nimic.
    Returnează lista blocurilor pre-completate (adăugată și în data/labels/jurnal_etichetare.csv)."""
    from openpyxl import load_workbook
    nom = pd.read_csv(INTERIM / "nomenclator_norm.csv", dtype=str).set_index("cod")
    ro = pd.read_csv(INTERIM / "cnam_ro_norm.csv", dtype=str)
    wb = load_workbook(path)
    ws = wb["etichetare"]
    col = {c.value: c.column_letter for c in ws[1]}
    blocks = {}
    for row in ws.iter_rows(min_row=2):
        blocks.setdefault(int(row[0].value), []).append(row[0].row)
    done = []
    for id_cnam, rows in blocks.items():
        cnam_row, cand_rows = rows[0], rows[1:]
        manual = [ws[f"{col['status']}{cnam_row}"].value, ws[f"{col['cod manual']}{cnam_row}"].value] + \
                 [ws[f"{col['alege (x)']}{r}"].value for r in cand_rows]
        if any(v not in (None, "") for v in manual):
            continue                                   # blocul are deja ceva completat de mână
        r = ro.loc[id_cnam]
        key = (r.denumire_norm, r.doza_norm, r.forma_norm, key_text(pd.Series([r.divizare])).iloc[0])
        same = []
        for i in cand_rows:
            cod = str(ws[f"{col['cod']}{i}"].value)
            c = nom.loc[cod] if cod in nom.index else None
            if isinstance(c, pd.DataFrame):
                c = c.iloc[0]
            if c is not None and (c.denumire_norm, c.doza_norm, c.forma_norm,
                                  key_text(pd.Series([c.divizare])).iloc[0]) == key:
                same.append((i, cod))
        if len(same) == 1:
            i, cod = same[0]
            ws[f"{col['alege (x)']}{i}"] = "x"
            ws[f"{col['comentariu']}{cnam_row}"] = PREFILL_NOTE
            done.append({"id_cnam": id_cnam, "cod": cod})
    wb.save(path)
    out = pd.DataFrame(done, columns=["id_cnam", "cod"])
    log_changes(out.assign(tip="pre-completare (candidat identic)", inainte="", dupa="x " + out.cod,
                           motiv="denumire, doză, formă și divizare identice"))
    return out


PREFILL_NOTE_UNITS = "x pre-completat automat (aceeași cantitate, divizare scrisă diferit), de verificat"
PREFILL_NOTE_MISSING = "produs există, divizare lipsă (pre-completat automat, de verificat)"


def _units(divizare) -> float:
    """Numărul total de unități din ambalaj; volumul scris înainte („60 ml N1”) este ignorat."""
    from clean import units_in_pack
    if not isinstance(divizare, str):
        return np.nan
    d = divizare.strip().lower()
    return units_in_pack(d[d.find("n"):] if "n" in d else d)


def _same_firm(detinator, cnam_text) -> bool:
    """Primul cuvânt semnificativ din numele deținătorului apare în textul CNAM (deținător + producător)."""
    words = [w for w in strip_accents(key_text(pd.Series([detinator]))).iloc[0].replace(",", " ").split()
             if len(w) > 2] if isinstance(detinator, str) else []
    text = strip_accents(key_text(pd.Series([cnam_text]))).iloc[0] if isinstance(cnam_text, str) else ""
    return bool(words) and words[0] in text


def prefill_divizare(path: Path = LABELS / "etichetare_cnam.xlsx") -> pd.DataFrame:
    """A doua pre-completare, pentru blocurile încă necompletate în care există candidați cu aceeași
    denumire, doză și formă normalizate și aceeași firmă, dar cu divizarea scrisă diferit:
    un singur ambalaj cu același număr de unități -> „x”; niciunul -> NICIUNUL + „produs există, divizare lipsă”.
    Blocurile cu mai multe ambalaje posibile rămân pentru etichetarea manuală. Nu șterge nimic."""
    from openpyxl import load_workbook
    nom = pd.read_csv(INTERIM / "nomenclator_norm.csv", dtype=str).drop_duplicates("cod").set_index("cod")
    ro = pd.read_csv(INTERIM / "cnam_ro_norm.csv", dtype=str)
    wb = load_workbook(path)
    ws = wb["etichetare"]
    col = {c.value: c.column_letter for c in ws[1]}
    blocks = {}
    for row in ws.iter_rows(min_row=2):
        blocks.setdefault(int(row[0].value), []).append(row[0].row)
    done = []
    for id_cnam, rows in blocks.items():
        cnam_row, cand_rows = rows[0], rows[1:]
        manual = [ws[f"{col[c]}{cnam_row}"].value for c in ("status", "cod manual", "comentariu")] + \
                 [ws[f"{col['alege (x)']}{r}"].value for r in cand_rows]
        if any(v not in (None, "") for v in manual):
            continue
        r = ro.loc[id_cnam]
        same = []
        for i in cand_rows:
            cod = str(ws[f"{col['cod']}{i}"].value)
            if cod not in nom.index:
                continue
            c = nom.loc[cod]
            if ((c.denumire_norm, c.doza_norm, c.forma_norm) == (r.denumire_norm, r.doza_norm, r.forma_norm)
                    and _same_firm(c.detinator, r.producator)):
                same.append((i, cod, _units(c.divizare)))
        if not same:
            continue
        u = _units(r.divizare)
        if np.isnan(u):
            continue
        hit = [(i, cod) for i, cod, uc in same if uc == u]
        if len(hit) == 1:
            ws[f"{col['alege (x)']}{hit[0][0]}"] = "x"
            ws[f"{col['comentariu']}{cnam_row}"] = PREFILL_NOTE_UNITS
            done.append({"id_cnam": id_cnam, "cod": hit[0][1], "regula": "aceeași cantitate"})
        elif not hit and not any(np.isnan(uc) for _, _, uc in same):
            ws[f"{col['status']}{cnam_row}"] = "NICIUNUL"
            ws[f"{col['comentariu']}{cnam_row}"] = PREFILL_NOTE_MISSING
            done.append({"id_cnam": id_cnam, "cod": None, "regula": "divizare lipsă"})
    wb.save(path)
    out = pd.DataFrame(done, columns=["id_cnam", "cod", "regula"])
    log_changes(out.assign(tip="pre-completare (divizare)", inainte="",
                           dupa=np.where(out.cod.notna(), "x " + out.cod.fillna(""), "NICIUNUL"), motiv=out.regula))
    return out


def add_volume_and_todo(path: Path = LABELS / "etichetare_cnam.xlsx") -> int:
    """Pentru fișierul generat înainte de adăugarea volumului: inserează coloana „volum” după „divizare”
    și adaugă coloana „de făcut” (DA pe toate rândurile blocurilor încă necompletate). Nu șterge nimic."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    nom = pd.read_csv(INTERIM / "nomenclator_norm.csv", dtype=str).drop_duplicates("cod").set_index("cod")
    ro = pd.read_csv(INTERIM / "cnam_ro_norm.csv", dtype=str)
    todo = set(read_labels(path).query("probleme == 'necompletat'").index)
    wb = load_workbook(path)
    ws = wb["etichetare"]
    head = [c.value for c in ws[1]]
    if "volum" not in head:
        pos = head.index("divizare") + 2                     # coloana imediat după „divizare” (1-based)
        ws.insert_cols(pos)
        ws.cell(1, pos, "volum")
        ws.cell(1, pos).font, ws.cell(1, pos).fill = Font(bold=True), PatternFill("solid", fgColor="D9D9D9")
        for i in range(2, ws.max_row + 1):
            tip, id_cnam, cod = ws.cell(i, 2).value, int(ws.cell(i, 1).value), str(ws.cell(i, head.index("cod") + 1).value)
            v = ro.loc[id_cnam, "volum"] if tip == "CNAM" else (nom.loc[cod, "volum"] if cod in nom.index else None)
            ws.cell(i, pos, v if isinstance(v, str) else None)
            ws.cell(i, pos).fill = ws.cell(i, pos - 1).fill.copy()
            ws.cell(i, pos).font = ws.cell(i, pos - 1).font.copy()
            ws.cell(i, pos).border = ws.cell(i, pos - 1).border.copy()
        for col, w in zip("MNO", [10, 28, 14]):
            ws.column_dimensions[col].width = w
    head = [c.value for c in ws[1]]
    pos = head.index("de făcut") + 1 if "de făcut" in head else ws.max_column + 1
    ws.cell(1, pos, "de făcut")
    ws.cell(1, pos).font, ws.cell(1, pos).fill = Font(bold=True), PatternFill("solid", fgColor="D9D9D9")
    for i in range(2, ws.max_row + 1):
        ws.cell(i, pos, "DA" if int(ws.cell(i, 1).value) in todo else None)
    ws.column_dimensions[ws.cell(1, pos).column_letter].width = 9
    ws.auto_filter.ref = ws.dimensions                       # filtre pe toate coloanele
    wb.save(path)
    return len(todo)


if __name__ == "__main__":
    LABELS.mkdir(parents=True, exist_ok=True)
    xlsx = LABELS / "etichetare_cnam.xlsx"
    if xlsx.exists():
        raise SystemExit(f"{xlsx} există deja; nu îl suprascriu (conține etichete manuale).")
    nom = pd.read_csv(INTERIM / "nomenclator_norm.csv", dtype=str)
    ro = pd.read_csv(INTERIM / "cnam_ro_norm.csv", dtype=str)
    nom["dci_k"], ro["dci_k"] = dci_key(nom.dci), dci_key(ro.dci)
    sample = draw_sample(ro, nom)
    sample[["id_cnam", "strat", "pondere"]].to_csv(LABELS / "esantion_info.csv", index=False)
    write_excel(sample, nom, xlsx)
    print(f"Eșantion: {len(sample)} rânduri CNAM ->", xlsx)
    print(sample.strat.value_counts().to_string())
